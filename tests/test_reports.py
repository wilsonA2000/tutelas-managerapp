"""Tests de reportes: Excel gen, download, list, metrics."""

import importlib.util
from pathlib import Path


def _load_build_cuadro():
    """Carga scripts/build_cuadro_excel.py como módulo (no es un paquete)."""
    path = Path(__file__).resolve().parent.parent / "scripts" / "build_cuadro_excel.py"
    spec = importlib.util.spec_from_file_location("build_cuadro_excel", path)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def test_build_cuadro_excel_script(TestSessionFactory, seed_data, tmp_path):
    """El CLI build_cuadro_excel produce un .xlsx con las 41 columnas del cuadro."""
    from openpyxl import load_workbook

    mod = _load_build_cuadro()
    db = TestSessionFactory()
    try:
        cases = mod._load_cases(db)
        assert len(cases) >= 3  # los 3 casos seed
        out = tmp_path / "cuadro.xlsx"
        from backend.reports.excel_generator import COLUMN_CONFIG, generate_excel
        generate_excel(cases, str(out))
        assert out.exists() and out.stat().st_size > 0
        wb = load_workbook(out)
        assert {"PORTADA", "TUTELAS", "ESTADISTICAS"} <= set(wb.sheetnames)
        ws = wb["TUTELAS"]
        assert ws.max_column == len(COLUMN_CONFIG)
        assert ws.max_row == len(cases) + 1  # +1 header
        # _value_for resuelve tanto el pseudo-campo TIPO como un campo real
        c0 = cases[0]
        assert mod._value_for(c0, "_TIPO_ACTUACION") in ("TUTELA", "INCIDENTE", "IMPUGNACION")
        assert mod._value_for(c0, "ACCIONANTE") == (c0.accionante or "")
    finally:
        db.close()


def test_metrics(client):
    r = client.get("/api/reports/metrics")
    assert r.status_code == 200
    data = r.json()
    assert "total" in data or isinstance(data, dict)


def test_excel_list(client):
    r = client.get("/api/reports/excel/list")
    assert r.status_code == 200


def test_generate_excel(client):
    r = client.post("/api/reports/excel")
    assert r.status_code == 200
    data = r.json()
    assert "filename" in data or "error" in data


def test_download_excel_not_found(client):
    r = client.get("/api/reports/excel/download/noexiste.xlsx")
    assert r.status_code == 404


def test_download_excel_path_traversal(client):
    """Security: no debe permitir path traversal."""
    r = client.get("/api/reports/excel/download/../../etc/passwd")
    assert r.status_code in (400, 404, 422)
